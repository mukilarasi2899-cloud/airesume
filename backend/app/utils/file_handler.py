"""
File Handling Utilities
Upload, validation, and processing of resume files
"""

import os
import aiofiles
from pathlib import Path
from typing import BinaryIO, Optional
from datetime import datetime
import hashlib
import PyPDF2
import docx
import pdfplumber

from app.config import settings
from app.core.exceptions import (
    FileProcessingError,
    InvalidFileTypeError,
    FileSizeExceededError
)
from app.core.logging import get_logger

logger = get_logger(__name__)


class FileHandler:
    """Handles file operations for resume uploads"""
    
    @staticmethod
    def validate_file_extension(filename: str) -> str:
        """
        Validate file extension
        
        Args:
            filename: Name of the file
            
        Returns:
            File extension (lowercase)
            
        Raises:
            InvalidFileTypeError: If extension is not allowed
        """
        ext = Path(filename).suffix.lower().lstrip('.')
        
        if ext not in settings.allowed_extensions_list:
            logger.warning(f"Invalid file extension: {ext}")
            raise InvalidFileTypeError(ext, settings.allowed_extensions_list)
        
        return ext
    
    @staticmethod
    def validate_file_size(content: bytes) -> int:
        """
        Validate file size
        
        Args:
            content: File content as bytes
            
        Returns:
            File size in bytes
            
        Raises:
            FileSizeExceededError: If file is too large
        """
        size = len(content)
        
        if size > settings.MAX_UPLOAD_SIZE:
            logger.warning(f"File size {size} exceeds limit {settings.MAX_UPLOAD_SIZE}")
            raise FileSizeExceededError(size, settings.MAX_UPLOAD_SIZE)
        
        return size
    
    @staticmethod
    def generate_unique_filename(original_filename: str, user_id: Optional[str] = None) -> str:
        """
        Generate unique filename using hash
        
        Args:
            original_filename: Original file name
            user_id: Optional user identifier
            
        Returns:
            Unique filename with original extension
        """
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        ext = Path(original_filename).suffix
        
        # Create hash from filename + timestamp + user_id
        hash_input = f"{original_filename}{timestamp}{user_id or ''}"
        file_hash = hashlib.md5(hash_input.encode()).hexdigest()[:8]
        
        return f"{timestamp}_{file_hash}{ext}"
    
    @staticmethod
    async def save_upload_file(
        file_content: bytes,
        filename: str,
        user_id: Optional[str] = None
    ) -> tuple[str, str]:
        """
        Save uploaded file to disk
        
        Args:
            file_content: File content as bytes
            filename: Original filename
            user_id: Optional user identifier
            
        Returns:
            Tuple of (file_path, unique_filename)
            
        Raises:
            FileProcessingError: If save operation fails
        """
        try:
            # Validate
            FileHandler.validate_file_extension(filename)
            FileHandler.validate_file_size(file_content)
            
            # Generate unique filename
            unique_filename = FileHandler.generate_unique_filename(filename, user_id)
            file_path = Path(settings.UPLOAD_DIR) / unique_filename
            
            # Save file
            async with aiofiles.open(file_path, 'wb') as f:
                await f.write(file_content)
            
            logger.info(f"File saved successfully: {unique_filename}")
            return str(file_path), unique_filename
            
        except (InvalidFileTypeError, FileSizeExceededError):
            raise
        except Exception as e:
            logger.error(f"Failed to save file: {str(e)}")
            raise FileProcessingError(f"Failed to save file: {str(e)}")
    
    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        """
        Extract text from PDF file
        
        Args:
            file_path: Path to PDF file
            
        Returns:
            Extracted text
            
        Raises:
            FileProcessingError: If extraction fails
        """
        try:
            text = ""
            
            # Try pdfplumber first (better for complex PDFs)
            try:
                with pdfplumber.open(file_path) as pdf:
                    for page in pdf.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + "\n"
            except Exception:
                # Fallback to PyPDF2
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    for page in pdf_reader.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + "\n"
            
            if not text.strip():
                raise FileProcessingError("No text could be extracted from PDF")
            
            logger.info(f"Extracted {len(text)} characters from PDF")
            return text.strip()
            
        except FileProcessingError:
            raise
        except Exception as e:
            logger.error(f"PDF extraction failed: {str(e)}")
            raise FileProcessingError(f"Failed to extract text from PDF: {str(e)}")
    
    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """
        Extract text from DOCX file
        
        Args:
            file_path: Path to DOCX file
            
        Returns:
            Extracted text
            
        Raises:
            FileProcessingError: If extraction fails
        """
        try:
            doc = docx.Document(file_path)
            
            # Extract text from paragraphs
            text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += "\n" + cell.text
            
            if not text.strip():
                raise FileProcessingError("No text could be extracted from DOCX")
            
            logger.info(f"Extracted {len(text)} characters from DOCX")
            return text.strip()
            
        except FileProcessingError:
            raise
        except Exception as e:
            logger.error(f"DOCX extraction failed: {str(e)}")
            raise FileProcessingError(f"Failed to extract text from DOCX: {str(e)}")
    
    @staticmethod
    def extract_text(file_path: str) -> str:
        """
        Extract text from file based on extension
        
        Args:
            file_path: Path to file
            
        Returns:
            Extracted text
            
        Raises:
            FileProcessingError: If extraction fails
        """
        ext = Path(file_path).suffix.lower()
        
        if ext == '.pdf':
            return FileHandler.extract_text_from_pdf(file_path)
        elif ext in ['.docx', '.doc']:
            return FileHandler.extract_text_from_docx(file_path)
        else:
            raise InvalidFileTypeError(ext, settings.allowed_extensions_list)
    
    @staticmethod
    def delete_file(file_path: str) -> bool:
        """
        Delete file from disk
        
        Args:
            file_path: Path to file
            
        Returns:
            True if deleted successfully
        """
        try:
            if os.path.exists(file_path):
                os.remove(file_path)
                logger.info(f"Deleted file: {file_path}")
                return True
            return False
        except Exception as e:
            logger.error(f"Failed to delete file {file_path}: {str(e)}")
            return False
